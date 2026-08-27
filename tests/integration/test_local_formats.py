# Path: tests/integration/test_local_formats.py
# Author: GHANMI Helmi
# Current Role: AI Engineer
# Past Role: Researcher in Applied Mathematics
# Research Profile: https://www.researchgate.net/profile/Ghanmi-Helmi

from pathlib import Path

import fitz
import openpyxl
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from document_parser import DocumentParser

pytestmark = pytest.mark.integration


def test_parse_text_and_csv(tmp_path: Path) -> None:
    text_path = tmp_path / "notes.txt"
    text_path.write_text("alpha\nbeta", encoding="utf-8")
    text_result = DocumentParser().parse(str(text_path))
    assert text_result["tool_used"] == "Built-in"
    assert "alpha" in text_result["content"]

    csv_path = tmp_path / "data.csv"
    csv_path.write_text("name,score\nAda,10\n", encoding="utf-8")
    csv_result = DocumentParser().parse(str(csv_path))
    assert csv_result["tool_used"] == "Built-in CSV"
    assert "Ada" in csv_result["content"]


def test_parse_office_formats(tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Senior engineering document")
    document.save(docx_path)
    assert "Senior engineering" in DocumentParser().parse(str(docx_path))["content"]

    pptx_path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "Architecture review"
    presentation.save(pptx_path)
    assert "Architecture review" in DocumentParser().parse(str(pptx_path))["content"]

    xlsx_path = tmp_path / "sample.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["metric", "value"])
    sheet.append(["quality", 99])
    workbook.save(xlsx_path)
    assert "quality" in DocumentParser().parse(str(xlsx_path))["content"]


def test_parse_pdf_with_pymupdf(tmp_path: Path) -> None:
    pdf_path = tmp_path / "sample.pdf"
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "Digital PDF content for deterministic integration testing. " * 3)
        document.save(pdf_path)

    result = DocumentParser(tool="pymupdf", detect_scanned=True).parse(str(pdf_path))
    assert result["tool_used"] == "PyMuPDF"
    assert result["metadata"]["page_count"] == 1
    assert result["pdf_analysis"]["has_text"] is True
