# Path: src/document_parser/parser.py
# Author: GHANMI Helmi
# Current Role: AI Engineer
# Past Role: Researcher in Applied Mathematics
# Research Profile: https://www.researchgate.net/profile/Ghanmi-Helmi

"""Core document parsing service with optional parser backends."""

from __future__ import annotations

from contextlib import suppress
import csv
import importlib.util
import io
import logging
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urlparse

from .config import ParserSettings
from .exceptions import FileTooLargeError, OptionalDependencyError, UnsupportedDocumentError
from .models import ParseResult
from .security import download_to_temp

logger = logging.getLogger(__name__)
ParserMethod = Callable[[str], ParseResult]


class DocumentParser:
    """Parse common document formats into a stable page-aware result contract."""

    SUPPORTED_EXTENSIONS: dict[str, tuple[str, ...]] = {
        "pdf": ("pdf",),
        "word": ("docx",),
        "powerpoint": ("pptx",),
        "excel": ("xlsx",),
        "csv": ("csv",),
        "text": ("txt", "md", "markdown", "html", "htm"),
    }

    TOOL_COMPATIBILITY: dict[str, set[str]] = {
        "pymupdf": {"pdf"},
        "pymupdf4llm": {"pdf"},
        "pdfplumber": {"pdf"},
        "tesseract_ocr": {"pdf"},
        "easyocr": {"pdf"},
        "docling": {"pdf"},
        "unstructured": {"pdf"},
        "python-docx": {"word"},
        "python-pptx": {"powerpoint"},
        "openpyxl": {"excel"},
        "builtin_csv": {"csv"},
        "builtin": {"text"},
    }

    def __init__(
        self,
        tool: str | None = None,
        detect_scanned: bool = True,
        extract_images: bool = False,
        ocr_lang: str = "eng",
        settings: ParserSettings | None = None,
    ) -> None:
        self.tool = tool
        self.detect_scanned = detect_scanned
        self.extract_images = extract_images
        self.ocr_lang = ocr_lang
        self.settings = settings or ParserSettings.from_env()

    def parse(self, document_path: str) -> ParseResult:
        """Parse a local file or validated remote URL and return normalized output."""
        if not document_path or not document_path.strip():
            raise ValueError("document_path must not be empty")

        source = document_path.strip()
        downloaded_path: Path | None = None
        source_name: str

        try:
            if source.startswith(("http://", "https://")):
                downloaded_path = download_to_temp(source, self.settings)
                path = downloaded_path
                source_name = Path(urlparse(source).path).name or path.name
            else:
                path = Path(source).expanduser().resolve()
                source_name = path.name

            if not path.exists() or not path.is_file():
                raise FileNotFoundError(f"Document not found: {source}")

            file_size = path.stat().st_size
            if file_size > self.settings.max_file_size_bytes:
                raise FileTooLargeError(
                    f"Document exceeds the {self.settings.max_file_size_mb} MB configured limit."
                )

            extension = path.suffix.lower().lstrip(".")
            document_type = self._get_document_type(extension)
            if document_type == "unknown":
                raise UnsupportedDocumentError(
                    f"Unsupported file extension '.{extension or '<none>'}'. "
                    f"Supported extensions: {', '.join(self.supported_extensions())}"
                )

            if self.tool:
                self._validate_tool_compatibility(self.tool, document_type)
                parser_method = self._get_parser_method(self.tool)
            else:
                parser_method = self._auto_select_parser(document_type, str(path))

            logger.info(
                "Parsing document source=%s type=%s parser=%s size_bytes=%s",
                source_name,
                document_type,
                parser_method.__name__,
                file_size,
            )
            result = parser_method(str(path))
            result["file_name"] = source_name
            result["file_type"] = document_type
            result["file_size"] = file_size
            return result
        finally:
            if downloaded_path is not None:
                with suppress(FileNotFoundError):
                    downloaded_path.unlink()

    @classmethod
    def supported_extensions(cls) -> list[str]:
        return sorted(f".{extension}" for extensions in cls.SUPPORTED_EXTENSIONS.values() for extension in extensions)

    def _get_document_type(self, extension: str) -> str:
        for document_type, extensions in self.SUPPORTED_EXTENSIONS.items():
            if extension in extensions:
                return document_type
        return "unknown"

    def _validate_tool_compatibility(self, tool_name: str, document_type: str) -> None:
        compatible_types = self.TOOL_COMPATIBILITY.get(tool_name)
        if compatible_types is None:
            raise UnsupportedDocumentError(f"Unknown parser tool: {tool_name}")
        if document_type not in compatible_types:
            raise UnsupportedDocumentError(
                f"Parser '{tool_name}' does not support document type '{document_type}'."
            )

    def _auto_select_parser(self, document_type: str, path: str) -> ParserMethod:
        if document_type == "pdf":
            if self.detect_scanned and self._is_scanned_pdf(path):
                if _modules_available("pytesseract", "pdf2image"):
                    return self._parse_with_tesseract
                if _modules_available("easyocr", "pdf2image"):
                    return self._parse_with_easyocr
                if _modules_available("docling"):
                    return self._parse_with_docling
                raise OptionalDependencyError(
                    "Scanned PDF detected but no OCR backend is installed. "
                    "Install the project with the 'ocr', 'easyocr', or 'docling' extra."
                )
            if _modules_available("pymupdf4llm"):
                return self._parse_with_pymupdf4llm
            return self._parse_with_pymupdf
        if document_type == "word":
            return self._parse_word
        if document_type == "powerpoint":
            return self._parse_powerpoint
        if document_type == "excel":
            return self._parse_excel
        if document_type == "csv":
            return self._parse_csv
        if document_type == "text":
            return self._parse_text
        raise UnsupportedDocumentError(f"Unsupported document type: {document_type}")

    def _get_parser_method(self, tool_name: str) -> ParserMethod:
        method_map: dict[str, ParserMethod] = {
            "pymupdf": self._parse_with_pymupdf,
            "pymupdf4llm": self._parse_with_pymupdf4llm,
            "docling": self._parse_with_docling,
            "unstructured": self._parse_with_unstructured,
            "pdfplumber": self._parse_with_pdfplumber,
            "python-docx": self._parse_word,
            "python-pptx": self._parse_powerpoint,
            "openpyxl": self._parse_excel,
            "builtin_csv": self._parse_csv,
            "builtin": self._parse_text,
            "tesseract_ocr": self._parse_with_tesseract,
            "easyocr": self._parse_with_easyocr,
        }
        try:
            return method_map[tool_name]
        except KeyError as exc:
            raise UnsupportedDocumentError(f"Unknown parser tool: {tool_name}") from exc

    def _is_scanned_pdf(self, pdf_path: str) -> bool:
        try:
            import fitz

            with fitz.open(pdf_path) as document:
                sample_count = min(self.settings.scanned_detection_sample_size, len(document))
                if sample_count == 0:
                    return False
                pages_without_text = sum(
                    1
                    for page_number in range(sample_count)
                    if len(document[page_number].get_text().strip()) < 50
                )
            return pages_without_text / sample_count >= self.settings.scanned_pdf_threshold
        except Exception:
            logger.exception("Failed to inspect PDF for scanned content")
            return False

    def _analyze_pdf(self, pdf_path: str) -> dict[str, Any]:
        try:
            import fitz

            scanned_pages: list[int] = []
            hybrid_pages: list[int] = []
            digital_pages: list[int] = []
            blank_pages: list[int] = []

            with fitz.open(pdf_path) as document:
                total_pages = len(document)
                for page_number, page in enumerate(document, start=1):
                    text = page.get_text().strip()
                    images = page.get_images()
                    has_text = len(text) >= 50
                    has_images = bool(images)
                    if has_text and has_images:
                        hybrid_pages.append(page_number)
                    elif has_text:
                        digital_pages.append(page_number)
                    elif has_images:
                        scanned_pages.append(page_number)
                    else:
                        blank_pages.append(page_number)

            if total_pages == 0:
                pdf_type = "Empty"
            elif len(blank_pages) == total_pages:
                pdf_type = "Blank"
            elif scanned_pages and not (digital_pages or hybrid_pages):
                pdf_type = "Scanned"
            elif digital_pages and not (scanned_pages or hybrid_pages):
                pdf_type = "Digital"
            else:
                pdf_type = "Hybrid"

            return {
                "type": pdf_type,
                "total_pages": total_pages,
                "scanned_pages": scanned_pages,
                "hybrid_pages": hybrid_pages,
                "digital_pages": digital_pages,
                "blank_pages": blank_pages,
                "has_text": bool(digital_pages or hybrid_pages),
            }
        except Exception:
            logger.exception("Failed to analyze PDF")
            return {}

    def _parse_with_pymupdf(self, pdf_path: str) -> ParseResult:
        import fitz

        content: list[str] = []
        pages: list[dict[str, Any]] = []
        with fitz.open(pdf_path) as document:
            for page_number, page in enumerate(document, start=1):
                text = page.get_text()
                pages.append(
                    {
                        "page_number": page_number,
                        "content": text,
                        "metadata": {"width": page.rect.width, "height": page.rect.height},
                    }
                )
                content.append(f"## Page {page_number}\n\n{text}\n")
            images = self._extract_images_pymupdf(document) if self.extract_images else []

        return {
            "tool_used": "PyMuPDF",
            "content": "\n".join(content),
            "pages": pages,
            "images": images,
            "metadata": {"page_count": len(pages)},
            "pdf_analysis": self._analyze_pdf(pdf_path) if self.detect_scanned else None,
        }

    def _parse_with_pymupdf4llm(self, pdf_path: str) -> ParseResult:
        try:
            import fitz
            import pymupdf4llm
        except ImportError as exc:
            raise OptionalDependencyError("PyMuPDF4LLM is not installed.") from exc

        markdown = pymupdf4llm.to_markdown(pdf_path)
        pages: list[dict[str, Any]] = []
        with fitz.open(pdf_path) as document:
            for page_number, page in enumerate(document, start=1):
                bboxes = []
                for block in page.get_text("dict").get("blocks", []):
                    if block.get("type") in {0, 1} and "bbox" in block:
                        bboxes.append(
                            {
                                "type": "text" if block["type"] == 0 else "image",
                                "bbox": block["bbox"],
                            }
                        )
                pages.append(
                    {
                        "page_number": page_number,
                        "content": page.get_text(),
                        "bboxes": bboxes,
                        "metadata": {"width": page.rect.width, "height": page.rect.height},
                    }
                )
            images = self._extract_images_pymupdf(document) if self.extract_images else []

        return {
            "tool_used": "PyMuPDF4LLM",
            "content": str(markdown),
            "pages": pages,
            "images": images,
            "metadata": {"page_count": len(pages)},
            "pdf_analysis": self._analyze_pdf(pdf_path) if self.detect_scanned else None,
        }

    def _parse_with_pdfplumber(self, pdf_path: str) -> ParseResult:
        try:
            import pdfplumber
        except ImportError as exc:
            raise OptionalDependencyError("pdfplumber is not installed.") from exc

        content: list[str] = []
        pages: list[dict[str, Any]] = []
        with pdfplumber.open(pdf_path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                tables = page.extract_tables()
                table_blocks = [_table_to_markdown(table) for table in tables if table]
                page_content = "\n\n".join(part for part in [text, *table_blocks] if part)
                pages.append(
                    {
                        "page_number": page_number,
                        "content": page_content,
                        "metadata": {
                            "width": page.width,
                            "height": page.height,
                            "tables_count": len(tables),
                        },
                    }
                )
                content.append(f"## Page {page_number}\n\n{page_content}\n")

        return {
            "tool_used": "pdfplumber",
            "content": "\n".join(content),
            "pages": pages,
            "images": [],
            "metadata": {"page_count": len(pages)},
            "pdf_analysis": self._analyze_pdf(pdf_path) if self.detect_scanned else None,
        }

    def _parse_with_tesseract(self, pdf_path: str) -> ParseResult:
        try:
            import pytesseract
            from pdf2image import convert_from_path
        except ImportError as exc:
            raise OptionalDependencyError(
                "Tesseract OCR requires the 'ocr' extra plus system packages tesseract-ocr and poppler-utils."
            ) from exc

        rendered_pages = convert_from_path(pdf_path, dpi=self.settings.pdf_render_dpi)
        pages: list[dict[str, Any]] = []
        content: list[str] = []
        for page_number, image in enumerate(rendered_pages, start=1):
            text = pytesseract.image_to_string(image, lang=self.ocr_lang)
            pages.append(
                {
                    "page_number": page_number,
                    "content": text,
                    "metadata": {"width": image.width, "height": image.height},
                }
            )
            content.append(f"## Page {page_number}\n\n{text}\n")
        return {
            "tool_used": "Tesseract OCR",
            "content": "\n".join(content),
            "pages": pages,
            "images": [],
            "metadata": {"page_count": len(pages)},
            "pdf_analysis": {"type": "Scanned", "total_pages": len(pages)},
        }

    def _parse_with_easyocr(self, pdf_path: str) -> ParseResult:
        try:
            import easyocr
            import numpy as np
            from pdf2image import convert_from_path
        except ImportError as exc:
            raise OptionalDependencyError("EasyOCR requires the 'easyocr' project extra.") from exc

        language = _easyocr_language(self.ocr_lang)
        reader = easyocr.Reader([language], gpu=False)
        rendered_pages = convert_from_path(pdf_path, dpi=self.settings.pdf_render_dpi)
        pages: list[dict[str, Any]] = []
        content: list[str] = []
        for page_number, image in enumerate(rendered_pages, start=1):
            detections = reader.readtext(np.array(image))
            text = "\n".join(str(item[1]) for item in detections)
            pages.append(
                {
                    "page_number": page_number,
                    "content": text,
                    "metadata": {
                        "detections": len(detections),
                        "width": image.width,
                        "height": image.height,
                    },
                }
            )
            content.append(f"## Page {page_number}\n\n{text}\n")
        return {
            "tool_used": "EasyOCR",
            "content": "\n".join(content),
            "pages": pages,
            "images": [],
            "metadata": {"page_count": len(pages)},
            "pdf_analysis": {"type": "Scanned", "total_pages": len(pages)},
        }

    def _parse_with_docling(self, document_path: str) -> ParseResult:
        try:
            from docling.document_converter import DocumentConverter
        except ImportError as exc:
            raise OptionalDependencyError("Docling requires the 'docling' project extra.") from exc

        try:
            result = DocumentConverter().convert(document_path, raises_on_error=False)
            if getattr(result, "status", None) and str(result.status).endswith("FAILURE"):
                errors = [str(error) for error in (result.errors or [])]
                return {
                    "tool_used": "Docling",
                    "content": "",
                    "pages": [],
                    "images": [],
                    "metadata": {},
                    "pdf_analysis": None,
                    "errors": errors or ["Docling conversion failed."],
                }

            document = result.document
            full_markdown = document.export_to_markdown(page_break_placeholder="\n\n---\n\n")
            page_count = len(document.pages)
            pages = [
                {
                    "page_number": page_number,
                    "content": document.export_to_markdown(page_no=page_number),
                    "metadata": {},
                }
                for page_number in range(1, page_count + 1)
            ]
            if not full_markdown.strip():
                full_markdown = document.export_to_text()
            return {
                "tool_used": "Docling",
                "content": full_markdown,
                "pages": pages,
                "images": [],
                "metadata": {"page_count": page_count},
                "pdf_analysis": None,
            }
        except Exception as exc:
            logger.exception("Docling conversion failed")
            raise RuntimeError("Docling conversion failed; see application logs for details.") from exc

    def _parse_with_unstructured(self, document_path: str) -> ParseResult:
        try:
            from unstructured.partition.auto import partition
        except ImportError as exc:
            raise OptionalDependencyError("Unstructured requires the 'unstructured' project extra.") from exc

        elements = partition(filename=document_path)
        content = "\n\n".join(str(element) for element in elements)
        return {
            "tool_used": "Unstructured",
            "content": content,
            "pages": [],
            "images": [],
            "metadata": {"elements_count": len(elements)},
            "pdf_analysis": None,
        }

    def _parse_word(self, document_path: str) -> ParseResult:
        try:
            from docx import Document
        except ImportError as exc:
            raise OptionalDependencyError("python-docx is not installed.") from exc

        document = Document(document_path)
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            blocks.append(_table_to_markdown(rows))
        content = "\n\n".join(blocks)
        return {
            "tool_used": "python-docx",
            "content": content,
            "pages": [],
            "images": [],
            "metadata": {"paragraphs": len(document.paragraphs), "tables": len(document.tables)},
        }

    def _parse_powerpoint(self, document_path: str) -> ParseResult:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise OptionalDependencyError("python-pptx is not installed.") from exc

        presentation = Presentation(document_path)
        pages: list[dict[str, Any]] = []
        content: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_text = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            slide_content = "\n".join(slide_text)
            pages.append({"page_number": slide_number, "content": slide_content, "metadata": {}})
            content.append(f"## Slide {slide_number}\n\n{slide_content}\n")
        return {
            "tool_used": "python-pptx",
            "content": "\n".join(content),
            "pages": pages,
            "images": [],
            "metadata": {"slides": len(presentation.slides)},
        }

    def _parse_excel(self, document_path: str) -> ParseResult:
        try:
            import openpyxl
        except ImportError as exc:
            raise OptionalDependencyError("openpyxl is not installed.") from exc

        workbook = openpyxl.load_workbook(document_path, data_only=True, read_only=True)
        try:
            blocks: list[str] = []
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows = [["" if cell is None else str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
                blocks.append(f"## Sheet: {sheet_name}\n\n{_table_to_markdown(rows)}")
            content = "\n\n".join(blocks)
            return {
                "tool_used": "openpyxl",
                "content": content,
                "pages": [],
                "images": [],
                "metadata": {"sheets": len(workbook.sheetnames)},
            }
        finally:
            workbook.close()

    def _parse_csv(self, document_path: str) -> ParseResult:
        with open(document_path, "r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            rows = list(csv.reader(handle))
        content = _table_to_markdown(rows)
        return {
            "tool_used": "Built-in CSV",
            "content": content,
            "pages": [],
            "images": [],
            "metadata": {"rows": len(rows)},
        }

    def _parse_text(self, document_path: str) -> ParseResult:
        with open(document_path, "r", encoding="utf-8", errors="replace") as handle:
            content = handle.read()
        return {
            "tool_used": "Built-in",
            "content": content,
            "pages": [],
            "images": [],
            "metadata": {"lines": len(content.splitlines())},
        }

    def _extract_images_pymupdf(self, document: Any) -> list[dict[str, Any]]:
        from PIL import Image

        images: list[dict[str, Any]] = []
        extracted_xrefs: set[int] = set()
        minimum_width = self.settings.min_image_width
        minimum_height = self.settings.min_image_height

        for page_number, page in enumerate(document, start=1):
            for image_index, image_info in enumerate(page.get_images(full=True)):
                xref = image_info[0]
                if xref in extracted_xrefs:
                    continue
                extracted_xrefs.add(xref)
                try:
                    raw = document.extract_image(xref)
                    image = Image.open(io.BytesIO(raw["image"]))
                    image.load()
                    if image.width < minimum_width or image.height < minimum_height:
                        continue
                    if image.mode not in {"RGB", "RGBA", "L"}:
                        image = image.convert("RGB")
                    images.append(
                        {
                            "page": page_number,
                            "index": image_index,
                            "image": image,
                            "ext": raw["ext"],
                            "width": image.width,
                            "height": image.height,
                            "mode": image.mode,
                        }
                    )
                except Exception:
                    logger.warning("Failed to extract image xref=%s page=%s", xref, page_number, exc_info=True)
        return images


def _modules_available(*module_names: str) -> bool:
    return all(importlib.util.find_spec(module_name) is not None for module_name in module_names)


def _escape_markdown_cell(value: Any) -> str:
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ")


def _table_to_markdown(rows: list[list[Any]]) -> str:
    if not rows:
        return ""
    width = max((len(row) for row in rows), default=0)
    if width == 0:
        return ""
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    lines = [
        "| " + " | ".join(_escape_markdown_cell(cell) for cell in header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(_escape_markdown_cell(cell) for cell in row) + " |" for row in normalized[1:])
    return "\n".join(lines)


def _easyocr_language(language: str) -> str:
    mapping = {
        "eng": "en",
        "fra": "fr",
        "deu": "de",
        "spa": "es",
        "ita": "it",
        "por": "pt",
        "rus": "ru",
        "chi_sim": "ch_sim",
        "chi_tra": "ch_tra",
        "jpn": "ja",
        "kor": "ko",
        "ara": "ar",
    }
    return mapping.get(language, language)
